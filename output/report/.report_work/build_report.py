from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\HCMUS\Ki_1_Nam_4\CDMMT\openguard-kids-project")
REFERENCE = ROOT / "output" / "report" / "23120357_23120049.docx"
OUT = ROOT / "output" / "report" / "BaoCao_OpenGuardKids_HoanChinh.docx"
WORK = ROOT / "output" / "report" / ".report_work"
ASSETS = WORK / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

EXPECTED_SHA = "B9F5290EF8CA4207461D8848CDCB3934D15B173685ED00A182DC2588EAB4C9E4"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def set_run_font(run, name="Times New Roman", size=13, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=11)


def add_toc_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở tài liệu trong Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)
    set_run_font(run, size=12)


def font(path_candidates, size):
    for p in path_candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


REGULAR = [r"C:\Windows\Fonts\times.ttf", r"C:\Windows\Fonts\arial.ttf"]
BOLD = [r"C:\Windows\Fonts\timesbd.ttf", r"C:\Windows\Fonts\arialbd.ttf"]


def draw_arrow(draw, start, end, fill="#4B5563", width=4):
    draw.line([start, end], fill=fill, width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    a = (x2 - 18 * ux + 8 * px, y2 - 18 * uy + 8 * py)
    b = (x2 - 18 * ux - 8 * px, y2 - 18 * uy - 8 * py)
    draw.polygon([end, a, b], fill=fill)


def rounded_box(draw, xy, title, lines, fill="#F5F7F6", outline="#325D50"):
    draw.rounded_rectangle(xy, radius=20, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = xy
    title_font = font(BOLD, 28)
    body_font = font(REGULAR, 22)
    draw.text(((x1 + x2) / 2, y1 + 28), title, font=title_font, fill="#122F27", anchor="ma")
    y = y1 + 78
    for line in lines:
        draw.text((x1 + 28, y), line, font=body_font, fill="#1F2937")
        y += 36


def make_architecture(path: Path):
    im = Image.new("RGB", (1600, 920), "white")
    d = ImageDraw.Draw(im)
    d.text((800, 45), "Kiến trúc triển khai OpenGuard Kids", font=font(BOLD, 40), fill="#111827", anchor="ma")
    rounded_box(d, (80, 180, 510, 720), "Máy của trẻ", [
        "Windows Service chạy LocalSystem",
        "Bộ đếm thời gian và lịch tuần",
        "Kiểm soát tiến trình và DNS proxy",
        "SQLite cache và event outbox",
        "DPAPI bảo vệ token và khóa",
        "Tray UI minh bạch với trẻ",
    ], fill="#ECF8F3")
    rounded_box(d, (610, 160, 1050, 740), "Server FastAPI", [
        "Xác thực phụ huynh và thiết bị",
        "Hồ sơ trẻ, policy và command",
        "WebSocket và heartbeat",
        "Sự kiện, báo cáo và audit",
        "Retention và xóa dữ liệu",
        "SQLite WAL và SQLAlchemy",
    ], fill="#F5F5F5", outline="#334155")
    rounded_box(d, (1140, 240, 1520, 660), "Dashboard phụ huynh", [
        "Quota và lịch sử dụng",
        "Luật web và ứng dụng",
        "Duyệt yêu cầu xin giờ",
        "Báo cáo và xuất PDF/CSV",
        "Theo dõi thiết bị, audit",
    ], fill="#EEF4FF", outline="#355A8A")
    draw_arrow(d, (510, 360), (610, 360), fill="#325D50")
    draw_arrow(d, (610, 470), (510, 470), fill="#325D50")
    d.text((560, 325), "HTTPS / WSS", font=font(BOLD, 20), fill="#325D50", anchor="mm")
    d.text((560, 510), "policy, event, command", font=font(REGULAR, 18), fill="#4B5563", anchor="mm")
    draw_arrow(d, (1050, 360), (1140, 360), fill="#355A8A")
    draw_arrow(d, (1140, 475), (1050, 475), fill="#355A8A")
    d.text((1095, 325), "HTTPS / SSE", font=font(BOLD, 20), fill="#355A8A", anchor="mm")
    d.text((1095, 515), "cấu hình và báo cáo", font=font(REGULAR, 18), fill="#4B5563", anchor="mm")
    d.text((800, 840), "Agent tiếp tục áp policy đã ký khi server tạm mất kết nối; dữ liệu được đồng bộ lại theo lô.", font=font(REGULAR, 23), fill="#374151", anchor="mm")
    im.save(path)


def make_dns_flow(path: Path):
    im = Image.new("RGB", (1600, 760), "white")
    d = ImageDraw.Draw(im)
    d.text((800, 40), "Luồng xử lý một truy vấn tên miền", font=font(BOLD, 38), fill="#111827", anchor="ma")
    boxes = [
        ((60, 220, 300, 480), "Trình duyệt", ["Gửi truy vấn A/AAAA"]),
        ((370, 220, 640, 480), "Stub resolver", ["Windows chuyển tới", "127.0.0.1:53"]),
        ((710, 160, 1080, 540), "DNS proxy của agent", ["Chuẩn hóa tên miền", "Tra cứu miễn chặn", "Tra danh sách riêng", "Tra nhóm phân loại", "Ghi domain_query"]),
        ((1160, 220, 1510, 480), "Phản hồi", ["NXDOMAIN nếu chặn", "SafeSearch alias", "hoặc chuyển upstream"]),
    ]
    for xy, title, lines in boxes:
        rounded_box(d, xy, title, lines, fill="#F7FAF9")
    for a, b in [((300, 350), (370, 350)), ((640, 350), (710, 350)), ((1080, 350), (1160, 350))]:
        draw_arrow(d, a, b, fill="#325D50")
    d.text((800, 650), "Proxy chỉ quan sát tên miền, không đọc URL, nội dung trang, tin nhắn hoặc dữ liệu nhập.", font=font(REGULAR, 25), fill="#374151", anchor="mm")
    im.save(path)


def make_sync_flow(path: Path):
    im = Image.new("RGB", (1600, 820), "white")
    d = ImageDraw.Draw(im)
    d.text((800, 35), "Vòng đời ghép đôi và đồng bộ chính sách", font=font(BOLD, 38), fill="#111827", anchor="ma")
    steps = [
        ("1", "Tạo mã ghép", "Phụ huynh xác nhận đồng ý; mã 8 ký tự, hạn 10 phút, dùng một lần."),
        ("2", "Enrollment", "Agent gửi mã và fingerprint; server trả device_id, access và refresh token."),
        ("3", "Đồng bộ", "Service mở WebSocket, gửi heartbeat và nhận policy theo phiên bản."),
        ("4", "Thực thi", "Agent kiểm HMAC, lưu cache, áp quota, lịch, ứng dụng và DNS."),
        ("5", "Phản hồi", "Event outbox gửi theo lô; server chống trùng và cập nhật báo cáo."),
    ]
    y = 150
    for idx, (n, title, detail) in enumerate(steps):
        x = 80 + idx * 300
        d.ellipse((x, y, x + 85, y + 85), fill="#325D50")
        d.text((x + 42, y + 42), n, font=font(BOLD, 32), fill="white", anchor="mm")
        d.text((x + 42, y + 125), title, font=font(BOLD, 24), fill="#122F27", anchor="ma")
        words = detail.split()
        lines, current = [], ""
        for word in words:
            trial = (current + " " + word).strip()
            if d.textlength(trial, font=font(REGULAR, 19)) > 250:
                lines.append(current)
                current = word
            else:
                current = trial
        if current:
            lines.append(current)
        yy = y + 175
        for line in lines:
            d.text((x - 80, yy), line, font=font(REGULAR, 19), fill="#374151")
            yy += 28
        if idx < len(steps) - 1:
            draw_arrow(d, (x + 95, y + 42), (x + 280, y + 42), fill="#6B7280", width=3)
    d.text((800, 735), "Khi mất mạng, policy đã ký và hàng đợi cục bộ giữ hệ thống hoạt động; reconnect sẽ tải lệnh còn chờ và đồng bộ sự kiện.", font=font(REGULAR, 22), fill="#374151", anchor="mm")
    im.save(path)


def crop_image(src: Path, dst: Path, box=None):
    im = Image.open(src).convert("RGB")
    if box:
        im = im.crop(box)
    im.save(dst, quality=94)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.75)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(5)

    for name, size, before, after in (
        ("Heading 1", 16, 12, 8),
        ("Heading 2", 14, 10, 5),
        ("Heading 3", 13, 8, 4),
    ):
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = False

    if "Caption" not in [s.name for s in styles]:
        styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    caption.font.size = Pt(11)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = True

    if "Code" not in [s.name for s in styles]:
        styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code = styles["Code"]
    code.font.name = "Consolas"
    code._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
    code._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(9.5)
    code.paragraph_format.left_indent = Cm(0.8)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.space_after = Pt(4)


def add_para(doc, text="", *, bold_lead=None, align=None, indent=True, italic=False, size=13, color="000000", keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size, italic=italic, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, italic=italic, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run("- " + item), size=12.5)
    return None


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("[Gợi ý ảnh cần bổ sung] " + text)
    set_run_font(r, size=11, italic=True, color="666666")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 13}[level], bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(text, style="Caption")
    for run in p.runs:
        set_run_font(run, size=11, italic=True)
    return p


def add_figure(doc, image_path: Path, caption: str, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    keep_row_together(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Cm(widths[i])
        set_cell_shading(cell, "3F4A4A")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.clear()
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cell, "F4F6F5")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.clear()
            r = p.add_run(str(text))
            set_run_font(r, size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_reference(doc, n, title, url):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(-0.7)
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"[{n}] {title}. {url}")
    set_run_font(r, size=11.5)


assert sha256(REFERENCE) == EXPECTED_SHA, "Template changed; distillation must be refreshed."
shutil.copy2(REFERENCE, OUT)

arch = ASSETS / "architecture.png"
dns = ASSETS / "dns-flow.png"
sync = ASSETS / "sync-flow.png"
time_shot = ASSETS / "time-policy.png"
filter_shot = ASSETS / "filter-dashboard.png"
report_shot = ASSETS / "reports.png"
make_architecture(arch)
make_dns_flow(dns)
make_sync_flow(sync)
crop_image(ROOT / "output" / "time-policy-fixed.png", time_shot)
crop_image(ROOT / "output" / "premium-dashboard-desktop.png", filter_shot, (0, 0, 1280, 2200))
crop_image(ROOT / "output" / "reports-analytics-final.png", report_shot)

doc = Document(OUT)
configure_styles(doc)

# Remove the template's old TOC content control, which is not exposed through
# Document.paragraphs and otherwise survives as broken PAGEREF fields.
body = doc._element.body
for child in list(body):
    if child.tag == qn("w:sdt"):
        instr = " ".join((node.text or "") for node in child.iter(qn("w:instrText")))
        if "TOC" in instr:
            body.remove(child)

# Remove the placeholder body but keep the complete cover through paragraph 18.
for p in list(doc.paragraphs)[19:]:
    parent = p._element.getparent()
    parent.remove(p._element)

# Keep the cover unnumbered and number the report body.
section = doc.sections[0]
section.different_first_page_header_footer = True
section.footer.is_linked_to_previous = False
footer_p = section.footer.paragraphs[0]
footer_p.clear()
add_page_field(footer_p)
section.first_page_footer.paragraphs[0].clear()

# Start the substantive report on a new page.
doc.add_page_break()

add_heading(doc, "TÓM TẮT", 1)
add_para(doc, "OpenGuard Kids là hệ thống hỗ trợ phụ huynh quản lý việc sử dụng máy tính của trẻ em trên Windows theo hướng minh bạch. Hệ thống gồm Windows Service và Tray UI trên máy trẻ, server FastAPI, cơ sở dữ liệu SQLite và dashboard web. Khác với phần mềm giám sát lén, sản phẩm chỉ ghi nhận siêu dữ liệu cần thiết như thời gian sử dụng, tên ứng dụng và tên miền; trẻ nhìn thấy trạng thái kiểm soát, lý do bị chặn, dữ liệu đã ghi nhận và có thể gửi yêu cầu xin thêm thời gian.")
add_para(doc, "Ở trạng thái hiện tại, repo đã hình thành một luồng đầu cuối tương đối hoàn chỉnh cho môi trường phát triển và máy lab: phụ huynh đăng nhập, tạo hồ sơ, ghép thiết bị bằng mã dùng một lần, cấu hình quota và lịch tuần, đẩy policy có phiên bản và HMAC, gửi lệnh khẩn qua WebSocket, nhận yêu cầu của trẻ, kiểm soát tiến trình và DNS, đồng bộ sự kiện, lập báo cáo ngày/tuần, xuất PDF hoặc CSV, ghi audit và xử lý xóa dữ liệu hoạt động. Các lớp bảo vệ chính gồm Argon2id, CSRF, kiểm tra quyền sở hữu, TLS 1.2+, ghim chứng chỉ, token xoay vòng, AES-GCM và DPAPI.")
add_para(doc, "Báo cáo này mô tả những gì đã được hiện thực trong mã nguồn, đối chiếu với yêu cầu F1-F5 và sáu mục kỹ thuật M1-M6, đồng thời chỉ rõ phần nào mới được kiểm chứng bằng kiểm thử tự động, phần nào cần nghiệm thu trên Windows thật. Kết luận chính là hệ thống đã đáp ứng phần lớn luồng chức năng ở mức development/lab, nhưng chưa nên xem là sản phẩm production vì còn thiếu đo hiệu năng đầy đủ, kiểm thử chống vượt rào trên máy thật, migration chuẩn, hạ tầng đa tiến trình và cơ chế chống can thiệp ở mức quản trị viên.")

add_heading(doc, "MỤC LỤC", 1)
toc_p = doc.add_paragraph()
toc_p.paragraph_format.first_line_indent = Cm(0)
add_toc_field(toc_p)

add_heading(doc, "DANH MỤC TỪ VIẾT TẮT", 1)
add_table(doc, ["Từ viết tắt", "Ý nghĩa"], [
    ("ACL", "Access Control List, danh sách quyền trên tệp và thư mục Windows"),
    ("API", "Application Programming Interface, giao diện lập trình ứng dụng"),
    ("CSRF", "Cross Site Request Forgery, giả mạo yêu cầu từ trình duyệt"),
    ("DNS", "Domain Name System, hệ thống phân giải tên miền"),
    ("DoH / DoT", "DNS over HTTPS / DNS over TLS"),
    ("DPAPI", "Data Protection API của Windows"),
    ("F1-F5", "Năm nhóm chức năng trong đề bài"),
    ("HMAC", "Mã xác thực thông điệp dựa trên hàm băm"),
    ("IPC", "Inter Process Communication, giao tiếp giữa Service và Tray"),
    ("SSE", "Server Sent Events, kênh cập nhật dashboard"),
    ("WAL", "Write Ahead Logging của SQLite"),
    ("WSS", "WebSocket Secure"),
], widths=[3.2, 12.4])

add_heading(doc, "1 GIỚI THIỆU ĐỀ TÀI", 1)
add_heading(doc, "1.1 Bối cảnh", 2)
add_para(doc, "Máy tính phục vụ học tập, giải trí và giao tiếp, nhưng trẻ nhỏ thường chưa tự phân bổ được thời gian hoặc nhận biết đầy đủ rủi ro trên Internet. Một công cụ hỗ trợ gia đình cần giúp phụ huynh đặt giới hạn hợp lý mà không biến máy tính thành thiết bị theo dõi bí mật. Từ yêu cầu đó, OpenGuard Kids được xây dựng như một hệ thống client-server: máy trẻ tự thực thi chính sách, còn phụ huynh cấu hình và xem báo cáo từ dashboard.")
add_para(doc, "Ràng buộc quan trọng nhất của đồ án là tính minh bạch. Agent không được ẩn, không đọc tin nhắn, tài liệu hay nội dung bàn phím, không chụp màn hình và không gửi dữ liệu sang dịch vụ bên thứ ba. Mỗi lần chặn phải có lý do; trẻ được xem thời gian còn lại, hoạt động đã ghi nhận, thay đổi chính sách và có kênh phản hồi. Cách tiếp cận này ảnh hưởng trực tiếp đến mô hình dữ liệu, giao diện, event schema và cả cách xử lý lỗi.")

add_heading(doc, "1.2 Mục tiêu và phạm vi", 2)
add_bullets(doc, [
    "Xây dựng kiến trúc phân tán gồm agent Windows, server điều phối và dashboard web.",
    "Kiểm soát thời gian sử dụng theo quota và lịch tuần, có cảnh báo và khoảng ân hạn.",
    "Kiểm soát ứng dụng theo tên tiến trình và SHA-256; lọc website qua DNS proxy cục bộ.",
    "Đồng bộ cấu hình từ xa, hỗ trợ lệnh khẩn, làm việc ngoại tuyến và chống trùng sự kiện.",
    "Cung cấp báo cáo cho phụ huynh và màn hình dễ hiểu cho trẻ trên cùng một tập dữ liệu.",
    "Áp dụng các kiểm soát bảo mật, quyền riêng tư và lưu vết phù hợp với phạm vi đồ án.",
])
add_para(doc, "Phạm vi triển khai là Windows 10/11 x64 trong môi trường phát triển hoặc máy lab. Server có thể chạy loopback hoặc trong LAN tin cậy; việc mở trực tiếp ra Internet không nằm trong phạm vi hiện tại. Sản phẩm không cam kết chống được người dùng có quyền Administrator, không giải mã HTTPS và không phân loại toàn bộ nội dung Internet theo thời gian thực.")

add_heading(doc, "1.3 Kết quả nổi bật của repo", 2)
add_table(doc, ["Nhóm", "Kết quả hiện có", "Mức kiểm chứng"], [
    ("F1", "Quota ngày thường/cuối tuần, lịch 7 x 48, idle/lock, cảnh báo 10-5-1 phút, grace 60 giây, chống lùi đồng hồ", "Unit và integration; cần acceptance dài giờ"),
    ("F2", "Chặn ứng dụng theo tên và SHA-256, DNS proxy UDP/TCP, phân nhóm, allow/block list, SafeSearch", "Tự động và bộ đo isolated; cần máy Windows thật"),
    ("F3", "Event outbox, báo cáo ngày/tuần, top 10, yêu cầu trẻ, PDF/CSV, retention và xóa", "Tự động; có ảnh dashboard"),
    ("F4", "Ghép đôi, heartbeat, WebSocket, policy version/HMAC, lệnh khẩn, audit, offline cache", "Unit, integration và mô phỏng"),
    ("F5", "Argon2id, session cookie, CSRF, ownership, TLS/pinning, token rotation, AES-GCM, DPAPI", "Kiểm thử bảo mật ở mức development"),
], widths=[1.6, 9.3, 5.0])

add_heading(doc, "2 YÊU CẦU VÀ NGUYÊN TẮC THIẾT KẾ", 1)
add_heading(doc, "2.1 Đối chiếu mục tiêu học tập", 2)
add_table(doc, ["Mục tiêu", "Cách repo đáp ứng"], [
    ("LO1", "Kiến trúc agent - server - dashboard, tài liệu API, heartbeat, WebSocket và policy version."),
    ("LO2", "DNS proxy 127.0.0.1:53, app controller theo process/SHA-256 và Windows Service."),
    ("LO3", "Threat model và tài liệu nêu DoH, VPN, đổi DNS, IP trực tiếp, Safe Mode, USB boot, tài khoản khác và kill service."),
    ("LO4", "Xác thực nhiều lớp, mã hóa bí mật, kiểm tra quyền sở hữu, audit và schema dữ liệu tối thiểu."),
    ("LO5", "Consent khi enrollment, retention 90 ngày, xóa dữ liệu, giao diện minh bạch và tài liệu PRIVACY."),
    ("LO6", "Bộ test rộng, checklist nghiệm thu và tài liệu chỉ rõ phần đã đo, chưa đo, giới hạn còn lại."),
], widths=[2.0, 13.6])

add_heading(doc, "2.2 Sáu nguyên tắc minh bạch", 2)
add_para(doc, "Sáu nguyên tắc của đề bài được chuyển thành hành vi cụ thể thay vì chỉ xuất hiện trong tài liệu. Tray phải hiện diện khi Service hoạt động; nếu không nhận heartbeat giao diện trong 15 giây, Service dừng và phục hồi cấu hình mạng do nó quản lý. Sự kiện chỉ giữ metadata. Trẻ có nút xin thêm giờ, xem phản hồi, xem audit gần nhất và nhận popup giải thích khi ứng dụng hoặc website bị chặn.")
add_table(doc, ["Nguyên tắc", "Biểu hiện trong mã nguồn và giao diện"], [
    ("NT1 Trẻ biết", "Tray UI, thông báo khởi động, tab giải thích ứng dụng ghi nhận gì."),
    ("NT2 Tối thiểu dữ liệu", "Không lưu URL đầy đủ, tiêu đề cửa sổ, nội dung phím, ảnh màn hình hay đường dẫn tệp."),
    ("NT3 Trẻ có tiếng nói", "Yêu cầu xin thêm giờ; lý do và trạng thái được đồng bộ hai phía."),
    ("NT4 Giải thích được", "Popup chặn có reason và rule_author; audit lưu giá trị cũ và mới."),
    ("NT5 Minh bạch hai chiều", "Dashboard và Tray dùng cùng dữ liệu báo cáo; trẻ xem thay đổi chính sách."),
    ("NT6 Ưu tiên an toàn", "Miền trường học/hỗ trợ khẩn cấp nằm trong danh sách miễn chặn; cấu hình mạng được phục hồi khi agent dừng."),
], widths=[4.0, 11.6])

add_heading(doc, "2.3 Khung pháp lý và chuẩn tham chiếu", 2)
add_para(doc, "Báo cáo sử dụng các văn bản được đề bài yêu cầu làm định hướng thiết kế. Luật Trẻ em 2016 đặt trọng tâm vào quyền riêng tư; Nghị định 13/2023/NĐ-CP dẫn đến cách ghi nhận sự đồng ý, giới hạn mục đích, tối thiểu dữ liệu, thời hạn lưu và xóa; Luật An ninh mạng 2018 và Quyết định 830/QĐ-TTg định hướng vai trò bảo vệ, hỗ trợ và giáo dục. OWASP ASVS được dùng như checklist bảo mật ứng dụng, còn NIST SP 800-63B hỗ trợ các quyết định về xác thực và vòng đời phiên.")
add_table(doc, ["Yêu cầu tham chiếu", "Quyết định thiết kế", "Bằng chứng trong repo"], [
    ("Quyền riêng tư của trẻ", "Không thu nội dung; giới hạn event schema", "docs/PRIVACY.md, event_schemas.py"),
    ("Đồng ý và mục đích xử lý", "Consent tại enrollment; mô tả mục đích từng trường", "routes.py, PRIVACY.md"),
    ("Hạn chế lưu trữ", "Retention 90 ngày; generation/ACK khi xóa", "event_routes.py, event_queue.py"),
    ("Xác thực an toàn", "Argon2id, khóa tạm, session 12 giờ, CSRF", "security.py, dependencies.py"),
    ("Kênh truyền", "TLS 1.2+, certificate pinning, bearer token", "secure_run.py, transport_security.py"),
], widths=[4.1, 6.0, 5.5])

add_heading(doc, "2.4 Những chức năng cố ý không triển khai", 2)
add_para(doc, "Repo không có keylogger, chụp màn hình, đọc tin nhắn, đọc tài liệu, theo dõi URL đầy đủ, kỹ thuật rootkit hay chế độ ẩn. DNS proxy chỉ nhìn thấy tên miền. Việc giới hạn thu thập làm giảm độ chi tiết của báo cáo, nhưng đó là đánh đổi có chủ đích để sản phẩm vẫn là công cụ đồng hành trong gia đình.")

add_heading(doc, "3 KIẾN TRÚC HỆ THỐNG", 1)
add_heading(doc, "3.1 Kiến trúc tổng thể", 2)
add_figure(doc, arch, "Hình 3.1 Kiến trúc tổng thể của hệ thống", width=6.35)
add_para(doc, "Windows Service là thành phần có đặc quyền để thực thi policy. Tray UI chạy trong phiên người dùng và giao tiếp với Service qua named pipe; cách tách này giúp giao diện luôn hiện với trẻ nhưng không trao quyền sửa policy cho tài khoản thường. Server FastAPI giữ danh tính, policy, command, event và audit. Dashboard được phục vụ cùng origin với API, giảm bề mặt cấu hình CORS và cho phép session cookie/CSRF hoạt động nhất quán.")

add_heading(doc, "3.2 Trách nhiệm các thành phần", 2)
add_table(doc, ["Thành phần", "Trách nhiệm", "Công nghệ chính"], [
    ("Windows Service", "Đồng bộ, thực thi F1/F2, quản lý DNS, hàng đợi sự kiện, IPC", "Python, pywin32, psutil, SQLite"),
    ("Tray UI", "Đăng nhập hồ sơ trẻ, thời gian còn lại, xin giờ, minh bạch và báo cáo", "Tkinter, named pipe"),
    ("Server", "Auth, enrollment, policy, command, report, audit, retention", "FastAPI, SQLAlchemy, Uvicorn"),
    ("Dashboard", "Cấu hình và theo dõi từ trình duyệt máy tính/điện thoại", "HTML, CSS, JavaScript"),
    ("CSDL server", "Lưu parent, child, device, policy, session, event, audit", "SQLite WAL"),
], widths=[3.3, 7.7, 4.6])

add_heading(doc, "3.3 Ghép đôi và đồng bộ", 2)
add_figure(doc, sync, "Hình 3.2 Vòng đời ghép đôi và đồng bộ chính sách", width=6.35)
add_para(doc, "Mã ghép gồm 8 ký tự, hết hạn sau 10 phút và chỉ dùng một lần. Khi tạo mã, phụ huynh đánh dấu đồng ý cho thiết bị xử lý dữ liệu của hồ sơ. Agent gửi mã cùng fingerprint đã băm; server trả device_id, access token ngắn hạn, refresh token dài hạn và vật liệu xác minh policy. Token được lưu bằng DPAPI LocalMachine trên Windows. Refresh token được xoay khi sử dụng và bị thu hồi khi phụ huynh gỡ thiết bị.")
add_para(doc, "Service duy trì WebSocket để nhận lệnh và thông báo policy_changed. Heartbeat là đường dự phòng: agent gửi trạng thái, hồ sơ đang hoạt động, số giây đã dùng và policy version; server trả version mới và command chưa xử lý. Nếu version lệch, agent tải document policy, kiểm HMAC rồi mới lưu cache và áp dụng.")

add_heading(doc, "3.4 Mô hình dữ liệu", 2)
add_para(doc, "Mô hình hiện tại hỗ trợ một phụ huynh có nhiều hồ sơ trẻ và mỗi hồ sơ có nhiều thiết bị. Policy thời gian và policy lọc gắn theo child_id; command và request mang cả device_id lẫn child_id để kiểm tra phạm vi. ActivityEvent dùng khóa ghép device_id/event_id để chống trùng. EventGeneration và EventDeletionAck tạo hàng rào xóa đối với thiết bị ngoại tuyến.")
add_table(doc, ["Nhóm bảng", "Bảng chính", "Vai trò"], [
    ("Danh tính", "parents, parent_sessions, admin_users, admin_sessions", "Xác thực và phiên"),
    ("Gia đình", "children, devices, enrollment_codes", "Hồ sơ, thiết bị và ghép đôi"),
    ("Chính sách", "policies, filtering_policies", "Quota, lịch, ứng dụng, website"),
    ("Điều khiển", "device_commands, time_requests, audit", "Lệnh, yêu cầu và lịch sử thay đổi"),
    ("Hoạt động", "activity_events, observed_apps", "Báo cáo và danh mục ứng dụng"),
    ("Xóa dữ liệu", "event_generations, event_deletion_acks", "Ngăn dữ liệu cũ quay lại sau khi xóa"),
], widths=[3.0, 6.8, 5.8])

add_heading(doc, "3.5 Chế độ ngoại tuyến và nhất quán", 2)
add_para(doc, "Khi server mất kết nối, agent tiếp tục dùng policy hợp lệ đã cache, duy trì bộ đếm thời gian và ghi sự kiện vào SQLite outbox. Event có UUID ổn định nên retry không làm báo cáo tăng gấp đôi. Command được lưu ở server trước khi đẩy qua WebSocket; nếu kênh trực tiếp lỗi, heartbeat hoặc reconnect sẽ nhận lại command đang chờ. Cách thiết kế này ưu tiên khả năng hoạt động liên tục nhưng vẫn giữ dấu vết để giải thích sau khi kết nối phục hồi.")

add_heading(doc, "4 HIỆN THỰC CÁC CHỨC NĂNG", 1)
add_heading(doc, "4.1 F1 Quản lý thời gian sử dụng", 2)
add_para(doc, "Policy thời gian gồm cờ bật/tắt, quota ngày thường, quota cuối tuần và lịch tuần 7 ngày. Mỗi ngày được mã hóa thành 48 ô nửa giờ. Dashboard cho phép thêm nhiều khoảng, tự gộp khoảng chồng nhau và kiểm tra lỗi trước khi gửi. Server cập nhật theo expected_version; nếu hai cửa sổ sửa cùng lúc, yêu cầu cũ bị từ chối thay vì âm thầm ghi đè.")
add_figure(doc, time_shot, "Hình 4.1 Giao diện cấu hình quota và lịch tuần", width=6.2)
add_para(doc, "Bộ đếm chỉ cộng active time khi phiên người dùng hợp lệ, máy không bị khóa và không idle quá 5 phút. Cảnh báo xuất hiện khi còn 10, 5 và 1 phút. Khi hết quota hoặc ra ngoài lịch, hệ thống cho 60 giây để trẻ lưu bài rồi gọi LockWorkStation. Khoản cộng giờ được lưu riêng cho ngày hiện tại, không reset số phút đã dùng và được áp idempotent theo command id.")
add_para(doc, "Để chống lùi đồng hồ, TimeController dùng time.monotonic() để tính delta và duy trì trusted wall time. Nếu giờ hệ thống lùi quá ngưỡng so với tiến trình đơn điệu, agent không tăng quota, bật trạng thái clock_tamper và ghi sự kiện. Đây là biện pháp phù hợp với phạm vi ứng dụng, dù thiết bị có quyền Administrator vẫn có thể can thiệp sâu hơn.")
add_note(doc, "Chụp Tray UI trên máy Windows thật ở ba trạng thái: còn 10 phút, đang grace 60 giây và đã khóa. Nên đặt ba ảnh cạnh nhau và ghi rõ policy thử nghiệm.")

add_heading(doc, "4.2 F2 Kiểm soát ứng dụng và website", 2)
add_para(doc, "App controller quan sát tiến trình trong phiên người dùng. Luật có thể dựa trên tên executable và SHA-256 của tệp, vì vậy đổi tên tệp không đủ để vượt rule đã gắn hash. Khi phát hiện ứng dụng bị chặn, Service kết thúc tiến trình và yêu cầu Tray hiển thị popup có tên ứng dụng, lý do và người đặt luật. Danh mục observed apps chỉ gửi tên ứng dụng, hash, thiết bị và lần thấy gần nhất; đường dẫn tệp không rời khỏi máy trẻ.")
add_para(doc, "DNS proxy lắng nghe UDP và TCP tại 127.0.0.1:53. Truy vấn được chuẩn hóa, so theo tên miền cụ thể nhất, toàn bộ tên miền con, danh sách miễn chặn, allow/block riêng và policy của sáu nhóm. Với truy vấn bị chặn, proxy trả NXDOMAIN; với Google/Bing khi SafeSearch bật, proxy trả alias cưỡng chế; các truy vấn còn lại được chuyển tới upstream đã cấu hình.")
add_figure(doc, dns, "Hình 4.2 Luồng xử lý truy vấn DNS", width=6.35)
add_figure(doc, filter_shot, "Hình 4.3 Dashboard quản lý nhóm website, SafeSearch và ứng dụng", width=5.35)
add_para(doc, "Bộ dữ liệu common/domain_categories.csv cung cấp phân loại ban đầu. Phụ huynh có thể thêm, sửa hoặc đưa một tên miền về trạng thái chưa phân loại cho từng hồ sơ. Bộ đo isolated trong repo ghi nhận 50/50 miền mô phỏng bị chặn, 20/20 miền giáo dục không bị chặn và 480 truy vấn UDP/TCP. Phép đo này xác nhận logic DNS trong môi trường cô lập, không chứng minh khả năng phân loại toàn bộ Internet.")
add_note(doc, "Chụp popup giải thích khi chặn một ứng dụng và một website trên máy lab. Ảnh cần thấy rõ lý do, tên luật và Tray đang hiện diện.")

add_heading(doc, "4.3 F3 Ghi nhận và báo cáo hoạt động", 2)
add_para(doc, "Event schema cốt lõi giữ đúng bảy trường của đề bài: ts, device_id, child_id, type, subject, duration_sec và policy_id. Metadata phụ như event_id, generation và explanation phục vụ chống trùng, xóa dữ liệu và giải thích hành động. Subject chỉ là tên ứng dụng hoặc tên miền cấp đăng ký, không phải URL đầy đủ. DNS query được gộp trong cửa sổ ngắn để A/AAAA và retry không làm số liệu tăng sai.")
add_table(doc, ["Loại sự kiện", "Ý nghĩa báo cáo"], [
    ("app_start / app_stop", "Thời gian chạy ứng dụng và top ứng dụng"),
    ("domain_query", "Tên miền được hệ điều hành hỏi; không đồng nghĩa với thời gian đọc trang"),
    ("blocked_app / blocked_domain", "Số lần chính sách bảo vệ can thiệp"),
    ("quota_warning / locked", "Giải thích cảnh báo và lần khóa"),
    ("unlock_request / override_granted", "Tương tác xin thêm giờ và quyết định của phụ huynh"),
], widths=[5.0, 10.6])
add_figure(doc, report_shot, "Hình 4.4 Báo cáo thời lượng, lượt chặn, top ứng dụng và tên miền", width=6.2)
add_para(doc, "Dashboard tổng hợp theo 1 ngày, 7 ngày và 90 ngày; hiển thị thời lượng theo ngày, xu hướng theo tuần, top 10 ứng dụng, top tên miền, lượt bị chặn và yêu cầu của trẻ. API có cả CSV UTF-8 BOM và PDF tuần. File đã tải nằm ngoài phạm vi tác vụ xóa trên server, vì vậy giao diện thông báo rõ phụ huynh phải tự quản lý tệp xuất.")
add_para(doc, "Retention chạy mỗi giờ ở server và theo vòng đồng bộ ở agent. Khi phụ huynh xóa dữ liệu hoạt động, server tăng generation rồi xóa event/request trong cùng transaction. Agent so generation, xóa history và outbox cũ, sau đó gửi ACK. Event của thế hệ cũ đến muộn bị loại bỏ, nhờ vậy thiết bị từng ngoại tuyến không thể làm dữ liệu đã xóa xuất hiện lại.")

add_heading(doc, "4.4 F4 Agent và cấu hình từ xa", 2)
add_para(doc, "Dashboard quản lý nhiều hồ sơ và thiết bị. Thay đổi policy được đẩy ngay qua WebSocket; heartbeat là dự phòng trong giới hạn 60 giây. Lệnh khóa hoặc cộng giờ được lưu bền vững trước khi gửi để vẫn đến agent sau reconnect. Audit ghi actor, thời điểm UTC, IP, action, old và new; Tray tải bản đã giới hạn theo device/child để trẻ xem những thay đổi liên quan.")
add_para(doc, "Named pipe là ranh giới giữa Service và Tray. Giao thức IPC giới hạn kích thước message, validate schema và gắn request id. Tray chạy singleton để tránh nhiều cửa sổ cạnh tranh. Service chỉ tiếp tục điều khiển khi có nhịp hiện diện của Tray, phù hợp với nguyên tắc không hoạt động ẩn.")
add_note(doc, "Chụp dashboard ở màn hình thiết bị online và một ảnh Tray hiển thị audit hoặc yêu cầu xin thêm giờ đã được duyệt.")

add_heading(doc, "4.5 F5 Bảo mật và quyền riêng tư", 2)
add_table(doc, ["Lớp bảo vệ", "Hiện thực"], [
    ("Mật khẩu và PIN", "Argon2id; tài khoản phụ huynh tối thiểu 10 ký tự; khóa 15 phút sau 5 lần sai"),
    ("Phiên web", "Cookie HttpOnly, Secure khi HTTPS, SameSite=Lax; hạn 12 giờ; logout xóa phiên; CSRF cho thao tác ghi"),
    ("Thiết bị", "Access 15 phút, refresh 30 ngày, xoay token, thu hồi theo device_id, fingerprint băm"),
    ("Kênh truyền", "TLS 1.2+; pin SHA-256 chứng chỉ trước HTTP/WebSocket; HTTP chỉ loopback hoặc ngoại lệ lab"),
    ("Dữ liệu lưu", "AES-256-GCM cho khóa policy phía server; DPAPI LocalMachine cho bí mật agent; token băm SHA-256"),
    ("Phân quyền", "Ownership ở tầng query; parent, child-view và admin development tách endpoint"),
    ("Toàn vẹn policy", "HMAC-SHA256 riêng theo thiết bị, version chống rollback"),
], widths=[4.2, 11.4])

add_heading(doc, "5 PHÂN TÍCH KỸ THUẬT HỆ THỐNG VÀ MẠNG", 1)
add_heading(doc, "5.1 M1 Luồng gói tin khi mở trang bị chặn", 2)
add_para(doc, "Trình duyệt gửi truy vấn cho Windows stub resolver. Khi Service đã cấu hình DNS của adapter về loopback, truy vấn đến proxy 127.0.0.1:53. Proxy chuẩn hóa hostname, tra miễn chặn, rule cụ thể, nhóm phân loại và SafeSearch. Nếu rule kết luận chặn, proxy trả NXDOMAIN và ghi blocked_domain; nếu cho phép, proxy chuyển tiếp tới upstream rồi trả kết quả cho stub resolver. Lưu đồ ở Hình 4.2 mô tả đường đi logic; bằng chứng packet capture vẫn cần lấy trên máy lab.")
add_note(doc, "Chèn ảnh Wireshark hoặc pktmon gồm một truy vấn UDP và một truy vấn TCP tới 127.0.0.1:53, kèm phản hồi NXDOMAIN. Che thông tin không liên quan và dùng tên miền thử nghiệm trung tính.")

add_heading(doc, "5.2 M2 So sánh phương án chặn", 2)
add_table(doc, ["Phương án", "Độ chính xác", "HTTPS", "Chi phí", "Khả năng bị vượt", "Ảnh hưởng"], [
    ("Tệp hosts", "Theo hostname", "Chặn được tên miền", "Thấp", "Dễ sửa/đổi DNS không liên quan", "Thấp"),
    ("DNS proxy", "Tên miền và tên miền con", "Không đọc đường dẫn", "Vừa", "DoH/DoT/VPN/IP trực tiếp", "Thấp-vừa"),
    ("HTTP(S) proxy", "Có thể theo URL nếu giải mã", "Cần MITM/chứng chỉ", "Cao", "VPN, pinning ứng dụng", "Vừa-cao"),
    ("WFP", "Theo luồng mạng", "Không tự hiểu nội dung", "Rất cao", "Khó vượt với user thường", "Tùy driver/rule"),
], widths=[2.7, 3.0, 3.0, 2.1, 3.1, 2.0], font_size=9.5)
add_para(doc, "Nhóm chọn DNS proxy vì phù hợp với yêu cầu chặn theo tên miền, không cần giải mã HTTPS và có thể giải thích rõ dữ liệu thu thập. Nhược điểm được giữ công khai: DNS không thấy URL, không ngăn được mọi ứng dụng tự dùng DoH hoặc VPN và không phải công cụ kiểm duyệt nội dung hoàn chỉnh.")

add_heading(doc, "5.3 M3 Các đường vòng và biện pháp", 2)
add_table(doc, ["Đường vòng", "Trạng thái hiện tại", "Biện pháp hoặc quyết định"], [
    ("Đổi DNS thủ công", "Agent có thể cấu hình và phục hồi DNS adapter; chưa nghiệm thu mọi adapter", "Tài khoản trẻ là standard user; kiểm tra cấu hình định kỳ"),
    ("DNS over HTTPS", "Secure DNS được cấu hình ở mức trình duyệt hỗ trợ; chưa có canary/blocklist đầy đủ", "Bổ sung endpoint list và policy trình duyệt; ghi nhận giới hạn"),
    ("DNS over TLS", "Chưa chặn ở tầng mạng", "Cần firewall/WFP trong bản nâng cấp"),
    ("VPN/proxy công cộng", "Có thể vượt DNS proxy", "Chấp nhận rủi ro ở bản lab; production cần WFP hoặc MDM"),
    ("Truy cập IP trực tiếp", "DNS không can thiệp", "Có thể chặn theo IP ở WFP; không ánh xạ mù vì CDN dùng chung"),
    ("Safe Mode", "Service có thể không chạy", "Giới hạn hệ điều hành; cần BIOS/BitLocker và quy trình gia đình"),
    ("Khởi động USB", "Ngoài phạm vi phần mềm", "Secure Boot, mật khẩu BIOS và mã hóa ổ đĩa"),
    ("Tài khoản Windows khác", "Chưa hỗ trợ multi-session hoàn chỉnh", "Giới hạn quyền tạo user; bổ sung session tracking"),
    ("Kết thúc tiến trình", "Service LocalSystem và SCM recovery; Tray mất thì Service tự dừng an toàn", "ACL và standard user; Administrator vẫn vượt được"),
], widths=[3.2, 6.1, 6.3], font_size=9.5)

add_heading(doc, "5.4 M4 Đo hiệu năng", 2)
add_para(doc, "Repo có bộ đo DNS isolated và các điểm ghi thời gian cho policy/command, nhưng chưa có bộ số liệu hoàn chỉnh gồm CPU, RAM, p50, p95 và tối thiểu 30 mẫu trên máy Windows thật. Vì vậy báo cáo không tự tạo số liệu giả. Kết quả hiện có chỉ nên dùng để xác nhận logic: 50/50 miền mô phỏng bị chặn, 20/20 miền giáo dục không bị chặn và 480 truy vấn UDP/TCP được thực hiện trong môi trường cô lập.")
add_table(doc, ["Phép đo cần bổ sung", "Cách đo đề xuất", "Kết quả cần ghi"], [
    ("CPU/RAM rảnh", "Performance Monitor hoặc psutil, lấy 30 mẫu cách nhau 2 giây", "Trung bình, p95, máy thử"),
    ("CPU/RAM cao điểm", "Phát truy vấn DNS và mở/đóng ứng dụng thử nghiệm", "Trung bình, p95"),
    ("DNS latency", "Cùng 50 hostname, đo có/không proxy", "p50, p95, chênh lệch"),
    ("Policy latency", "Timestamp lúc bấm lưu và lúc Service áp version", "30 mẫu, p50, p95, max"),
    ("Command latency", "Timestamp API và popup/LockWorkStation", "30 mẫu, tỷ lệ dưới 5 giây"),
], widths=[4.0, 7.0, 4.6])
add_note(doc, "Chèn bảng kết quả đo thật và ảnh Performance Monitor. Ghi cấu hình máy, phiên bản Windows, số mẫu, phương pháp và thời điểm đo.")

add_heading(doc, "5.5 M5 Khả năng chịu lỗi", 2)
add_table(doc, ["Kịch bản", "Hành vi thiết kế", "Tình trạng"], [
    ("Mất mạng", "Dùng policy cache, tiếp tục F1/F2, event vào outbox, reconnect theo heartbeat", "Có mã và test; cần demo thời gian phục hồi"),
    ("Server dừng", "WebSocket ngắt, command giữ ở DB, agent không bỏ policy hợp lệ", "Có test mô phỏng"),
    ("Agent bị kết thúc", "SCM recovery khởi động lại; state quota lưu SQLite; Tray mất thì dừng an toàn", "Có logic; cần nghiệm thu Windows"),
], widths=[3.4, 8.1, 4.1])

add_heading(doc, "5.6 M6 Đồng bộ thời gian", 2)
add_para(doc, "TimeController tách đồng hồ đơn điệu và thời gian tường. Số giây sử dụng được cộng theo delta monotonic; trusted wall time chỉ tiến, không lùi theo thao tác đổi giờ của người dùng. Heartbeat chuyển timestamp trạng thái cho server và audit dùng UTC. Khi phát hiện rollback, agent tạo event để phụ huynh biết thay vì lặng lẽ thay đổi quota.")

add_heading(doc, "6 BẢO MẬT QUYỀN RIÊNG VÀ MÔ HÌNH ĐE DỌA", 1)
add_heading(doc, "6.1 Dòng dữ liệu tối thiểu", 2)
add_para(doc, "Mỗi trường dữ liệu trong docs/PRIVACY.md đều có mục đích, thời hạn và người được xem. Đối với hoạt động, cửa sổ lưu là 90 ngày. Report chỉ tổng hợp event còn trong cửa sổ đó. Dữ liệu định danh và policy có vòng đời khác vì cần để tài khoản, thiết bị và rule tiếp tục hoạt động; giao diện tách rõ xóa lịch sử hoạt động với xóa toàn bộ tài khoản.")
add_para(doc, "SQLite secure_delete hỗ trợ xóa logic trên máy, nhưng dự án không tuyên bố xóa forensic khỏi WAL, snapshot, bản sao lưu hoặc tệp PDF đã tải. Giới hạn này được nói rõ trong tài liệu thay vì tạo cảm giác xóa tuyệt đối.")

add_heading(doc, "6.2 Mô hình STRIDE", 2)
add_table(doc, ["Nhóm", "Mối đe dọa tiêu biểu", "Kiểm soát hiện có", "Khoảng trống"], [
    ("Spoofing", "Đoán mã ghép hoặc giả thiết bị", "Mã ngẫu nhiên, hết hạn, một lần; token gắn device", "Rate limiter chưa phân tán"),
    ("Tampering", "Sửa policy/cache", "Version và HMAC; khóa mã hóa", "HMAC đối xứng"),
    ("Repudiation", "Phủ nhận thay đổi policy", "Audit actor, IP, old/new", "Audit chưa bất biến"),
    ("Information disclosure", "Lộ token/dữ liệu trẻ", "Hash token, AES-GCM, DPAPI, ownership", "SQLite không mã hóa toàn bộ"),
    ("Denial of service", "Làm server hoặc DNS proxy quá tải", "Rate limit và giới hạn IPC/batch", "Một process, chưa có queue ngoài"),
    ("Elevation", "Trẻ kill agent, đổi DNS, dùng admin", "Service, ACL, standard user, SCM recovery", "Administrator vẫn có thể vượt"),
    ("Abuse", "Dùng sản phẩm theo dõi người lớn", "Không chế độ ẩn, yêu cầu Tray, phạm vi sử dụng hợp pháp", "Cần installer và consent mạnh hơn"),
], widths=[2.2, 4.5, 5.1, 3.8], font_size=9.3)

add_heading(doc, "6.3 Checklist OWASP ASVS", 2)
add_para(doc, "Tài liệu F5 liệt kê 20 mục ASVS mức 1 liên quan đến kiến trúc, xác thực, phiên, kiểm soát truy cập, validation, mã hóa, logging, dữ liệu và truyền thông. Điểm quan trọng là ownership được kiểm tra ở tầng truy vấn, không chỉ ẩn nút trên giao diện; các thao tác ghi cần CSRF; WebSocket không đặt bearer token trong URL; lỗi ghim chứng chỉ làm kết nối thất bại thay vì hạ xuống HTTP.")

add_heading(doc, "7 KIỂM THỬ VÀ ĐÁNH GIÁ", 1)
add_heading(doc, "7.1 Cấu trúc kiểm thử", 2)
add_para(doc, "Thư mục tests hiện có 212 hàm test trong 15 tệp Python, bao phủ time controller, IPC, setup agent, child profile, filtering, report, remote sync, security, integration, realtime, Tray UI và web protection. Một test có thể được tham số hóa thành nhiều case; vì vậy số test pytest thực thi lớn hơn số hàm. README tại mốc hiện tại ghi nhận 266 passed. Lần chạy lại trong môi trường soạn báo cáo không hoàn tất vì một nhóm kiểm thử treo sau khi đã chạy một phần, nên số 266 được trình bày là kết quả đã ghi nhận trong repo, không phải kết quả tái xác nhận của phiên soạn tài liệu này.")
add_table(doc, ["Nhóm test", "Nội dung chính"], [
    ("test_time_control.py", "Quota, lịch, idle, cảnh báo, grace, rollback clock, persistence"),
    ("test_f2_filtering.py", "Rule tên miền, subdomain, SafeSearch, DNS UDP/TCP, app hash"),
    ("test_f3_reports.py", "Tổng hợp báo cáo, CSV/PDF, retention, quyền sở hữu"),
    ("test_f4_remote.py", "Policy version/HMAC, WebSocket, command, offline cache"),
    ("test_f5_security.py", "TLS/pinning, DPAPI/AES-GCM, secret handling"),
    ("test_integration.py", "Đăng nhập, enrollment, heartbeat, policy và device lifecycle"),
    ("test_tray_ui.py", "Màn hình minh bạch, request, popup và state UI"),
    ("test_web_protection.py", "Cấu hình/khôi phục DNS và Secure DNS"),
], widths=[5.0, 10.6])
add_note(doc, "Trước khi nộp, chạy lại `pytest -q`, `ruff check .` và `ruff format --check .`; chèn ảnh màn hình kết quả cuối cùng có thời gian chạy và số case pass.")

add_heading(doc, "7.2 Kiểm thử DNS an toàn", 2)
add_para(doc, "Bộ dữ liệu kiểm thử không chứa miền người lớn thật. Các miền cần chặn dùng example/test hoặc miền mô phỏng trung tính; nhóm giáo dục dùng địa chỉ công khai để đo dương tính giả. Script check_f2_dns.py chạy ở chế độ isolated, không đổi DNS adapter, không mở website và không sửa policy đang dùng. Thiết kế này vừa đáp ứng mục tiêu kỹ thuật vừa tránh đưa nội dung không phù hợp vào môi trường học tập.")

add_heading(doc, "7.3 Những gì kiểm thử tự động chưa chứng minh", 2)
add_bullets(doc, [
    "LockWorkStation hoạt động đúng trong phiên Windows thật và độ trễ luôn dưới 5 giây.",
    "Sai số quota sau một giờ sử dụng thực tế nằm trong khoảng cộng trừ 30 giây.",
    "DNS adapter và Secure DNS được phục hồi đúng trên mọi cấu hình mạng.",
    "Mức CPU/RAM và latency p50/p95 đạt mục tiêu trên máy cấu hình thấp.",
    "Installer chạy được trên máy Windows sạch và gỡ cài đặt hợp pháp không để lại service.",
])

add_heading(doc, "8 TRIỂN KHAI VÀ KỊCH BẢN DEMO", 1)
add_heading(doc, "8.1 Khởi động server", 2)
add_para(doc, "Quy trình phát triển dùng Python 3.11 trở lên, virtual environment và requirements-lock.txt. Sau khi tạo .env và khóa mã hóa, chạy seed rồi khởi động Uvicorn. Dashboard ở cùng server; tài khoản demo lấy từ .env.example và phải đổi trước lần seed đầu nếu muốn dùng thông tin khác.")
code_lines = [
    "python -m venv .venv",
    ".\\.venv\\Scripts\\python.exe -m pip install -r requirements-lock.txt",
    "Copy-Item .env.example .env",
    "& .\\scripts\\Initialize-OGKSecrets.ps1",
    ".\\.venv\\Scripts\\python.exe -m server.seed",
    ".\\.venv\\Scripts\\python.exe -m uvicorn server.app.main:create_app --factory --host 127.0.0.1 --port 8000",
]
for line in code_lines:
    p = doc.add_paragraph(style="Code")
    r = p.add_run(line)
    set_run_font(r, name="Consolas", size=9.5)

add_heading(doc, "8.2 Cài agent trên máy lab", 2)
add_para(doc, "Trên dashboard, phụ huynh tạo hồ sơ trẻ và mã ghép. Trên máy trẻ, chạy scripts/Setup-OGKAgent.ps1, chấp nhận UAC và nhập mã trong Tray. Nếu server ở máy khác, phải truyền URL HTTPS và SHA-256 certificate pin; HTTP LAN chỉ được phép khi bật cờ ngoại lệ lab. Tài khoản của trẻ phải là standard user.")
add_note(doc, "Chụp quá trình cài service thành công, Tray ghép đôi và dashboard chuyển thiết bị sang online. Không để lộ token, khóa hoặc nội dung file .env.")

add_heading(doc, "8.3 Kịch bản demo đề xuất", 2)
add_table(doc, ["Bước", "Thao tác", "Kết quả quan sát"], [
    ("1", "Đăng nhập, tạo hồ sơ và mã ghép", "Mã 8 ký tự, hạn 10 phút, có consent"),
    ("2", "Ghép agent và đăng nhập PIN trên Tray", "Thiết bị online, đúng hồ sơ"),
    ("3", "Đặt quota 3 phút và lịch đang cho phép", "Policy version tăng, agent tải policy"),
    ("4", "Chạy ứng dụng bị chặn", "Process bị dừng, popup giải thích, event xuất hiện"),
    ("5", "Truy vấn miền mô phỏng bị chặn", "NXDOMAIN, blocked_domain và báo cáo tăng"),
    ("6", "Trẻ xin 15 phút, phụ huynh duyệt", "Tray nhận kết quả, allowance tăng"),
    ("7", "Dừng server rồi sử dụng tiếp", "F1/F2 tiếp tục, event nằm trong outbox"),
    ("8", "Bật server và xóa dữ liệu hoạt động", "Outbox đồng bộ, sau đó generation/ACK xác nhận xóa"),
], widths=[1.5, 7.1, 7.0])

add_heading(doc, "9 ĐỐI CHIẾU MỨC ĐỘ HOÀN THÀNH", 1)
add_heading(doc, "9.1 Ma trận yêu cầu chức năng", 2)
add_table(doc, ["Yêu cầu", "Trạng thái", "Ghi chú trung thực"], [
    ("F1 quota và lịch", "Đã hiện thực", "Cần đo sai số một giờ và demo khóa thật"),
    ("F1 cảnh báo/grace", "Đã hiện thực", "Có test state machine; cần ảnh Windows"),
    ("F2 app name + SHA-256", "Đã hiện thực", "Cần nghiệm thu đổi tên file thật"),
    ("F2 DNS + subdomain", "Đã hiện thực", "Có isolated test; cần packet capture"),
    ("F2 6 nhóm và SafeSearch", "Đã hiện thực", "Dữ liệu seed cần rà soát thủ công"),
    ("F2 đối phó DoH", "Một phần", "Có cấu hình Secure DNS; chưa chặn mọi endpoint"),
    ("F3 report và child view", "Đã hiện thực", "Dashboard/Tray dùng cùng dữ liệu"),
    ("F3 retention/xóa", "Đã hiện thực", "Có generation/ACK; không cam kết forensic erase"),
    ("F3 xuất PDF/CSV", "Đã hiện thực", "Tệp tải về nằm ngoài tác vụ xóa"),
    ("F4 cấu hình dưới 60 giây", "Đã hiện thực về cơ chế", "Cần 30 mẫu latency đầu-cuối"),
    ("F4 lệnh khẩn dưới 5 giây", "Đã hiện thực về cơ chế", "Cần acceptance Windows"),
    ("F4 version/HMAC/audit", "Đã hiện thực", "HMAC đối xứng, audit chưa bất biến"),
    ("F5 bảo mật nền", "Đã hiện thực development", "Chưa có MFA production và limiter phân tán"),
], widths=[5.0, 3.2, 7.4], font_size=9.6)

add_heading(doc, "9.2 Các hạng mục M1-M6", 2)
add_table(doc, ["Hạng mục", "Mức hoàn thành", "Phần cần bổ sung trước khi nộp"], [
    ("M1", "Có sơ đồ và phân tích luồng", "Ảnh Wireshark/tcpdump thật"),
    ("M2", "Có bảng so sánh và lý do chọn DNS", "Không bắt buộc thêm"),
    ("M3", "Có threat model và 9 đường vòng", "Ghi kết quả thử thật từng kỹ thuật"),
    ("M4", "Có công cụ và kế hoạch đo", "CPU, RAM, p50, p95, policy latency, 30 mẫu"),
    ("M5", "Có thiết kế offline/reconnect/recovery", "Đo thời gian phục hồi trên Windows"),
    ("M6", "Đã dùng monotonic/trusted time", "Ảnh hoặc log thử lùi giờ 3 tiếng"),
], widths=[2.0, 6.2, 7.4])

add_heading(doc, "10 GIỚI HẠN VÀ HƯỚNG PHÁT TRIỂN", 1)
add_heading(doc, "10.1 Giới hạn hiện tại", 2)
add_bullets(doc, [
    "Server SQLite và registry realtime chạy một process; chưa phù hợp triển khai nhiều instance.",
    "Schema dùng create_all và migration idempotent tự viết; chưa có Alembic cho nâng cấp production.",
    "Rate limit ở bộ nhớ tiến trình, chưa chia sẻ giữa nhiều worker.",
    "HMAC policy dùng khóa đối xứng; chưa có chữ ký bất đối xứng và quy trình xoay khóa production.",
    "Agent chưa xử lý đầy đủ multi-session/RDP và không chống được tài khoản Administrator.",
    "DoH/DoT, VPN, IP trực tiếp, Safe Mode và USB boot vẫn có thể vượt một phần cơ chế DNS.",
    "Chưa có installer được ký số, auto-update, telemetry vận hành tối thiểu và rollback phiên bản.",
])

add_heading(doc, "10.2 Hướng phát triển", 2)
add_para(doc, "Bước tiếp theo nên ưu tiên hoàn tất bằng chứng nghiệm thu trước khi mở rộng chức năng. Cụ thể, nhóm cần chạy bộ đo M4, packet capture M1, kiểm thử M3 trên máy lab và quay video demo đầu cuối. Sau đó mới chuyển sang Alembic, PostgreSQL, Redis-backed rate limit/realtime, worker bền vững, chữ ký Ed25519, installer có chữ ký và cơ chế cập nhật. Nếu cần chống bypass mạnh hơn, WFP hoặc chính sách thiết bị tập trung phù hợp hơn việc tiếp tục chồng rule lên DNS proxy.")

add_heading(doc, "11 KẾT LUẬN", 1)
add_para(doc, "OpenGuard Kids đã đi từ bộ khung giao tiếp thành một hệ thống development có đủ các thành phần chính của đề bài. Luồng phụ huynh - server - agent hoạt động thống nhất từ ghép đôi, cấu hình, thực thi đến báo cáo. Điểm đáng chú ý không nằm ở số lượng chức năng, mà ở việc các chức năng dùng chung nguyên tắc: dữ liệu tối thiểu, policy có thể kiểm tra, hành động có lời giải thích và trẻ nhìn thấy hệ thống đang làm gì.")
add_para(doc, "Repo hiện phù hợp để trình diễn và bảo vệ kiến trúc trong môi trường lab. Nhóm đã có nền kiểm thử tốt và tài liệu nêu giới hạn tương đối rõ. Để kết luận đáp ứng trọn vẹn đề bài, phần còn thiếu chủ yếu là bằng chứng thực nghiệm trên Windows thật: latency, tài nguyên, sai số một giờ, packet capture, thử bypass và quy trình cài/gỡ trên máy sạch. Báo cáo giữ các khoảng trống này công khai để việc đánh giá phản ánh đúng trạng thái sản phẩm.")

add_heading(doc, "PHÂN CÔNG VÀ TỰ ĐÁNH GIÁ", 1)
add_para(doc, "Phần phân công chưa thể suy ra chính xác từ lịch sử repo và cần hai thành viên xác nhận trước khi nộp. Nhóm nên điền công việc thực tế, bằng chứng commit hoặc tài liệu phụ trách và tỷ lệ đóng góp sao cho tổng bằng 100 phần trăm.")
add_table(doc, ["Sinh viên", "MSSV", "Phần việc", "Tỷ lệ"], [
    ("Lê Nhật Thành", "23120357", "[Nhóm bổ sung]", "[ ] %"),
    ("Nguyễn Thanh Huyền", "23120049", "[Nhóm bổ sung]", "[ ] %"),
], widths=[4.2, 3.0, 6.4, 2.0])

add_heading(doc, "DANH MỤC ẢNH CẦN NHÓM BỔ SUNG", 1)
add_table(doc, ["Vị trí", "Ảnh cần chụp", "Mục đích"], [
    ("Mục 4.1", "Tray ở cảnh báo, grace và khóa", "Chứng minh F1 trên Windows thật"),
    ("Mục 4.2", "Popup chặn app và website", "Chứng minh NT4 và F2"),
    ("Mục 4.4", "Thiết bị online, audit và xin giờ", "Chứng minh đồng bộ hai chiều"),
    ("Mục 5.1", "Wireshark/pktmon UDP và TCP DNS", "Bằng chứng M1"),
    ("Mục 5.4", "Performance Monitor và bảng 30 mẫu", "Bằng chứng M4"),
    ("Mục 7.1", "Kết quả pytest/ruff cuối cùng", "Bằng chứng kiểm thử"),
    ("Mục 8.2", "Cài Service, ghép Tray, dashboard online", "Bằng chứng cài đặt"),
], widths=[2.4, 7.0, 6.2])

add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
refs = [
    ("Quốc hội, Luật Trẻ em số 102/2016/QH13", "https://vanban.chinhphu.vn/?docid=184566&pageid=27160"),
    ("Chính phủ, Nghị định 13/2023/NĐ-CP về bảo vệ dữ liệu cá nhân", "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=161106"),
    ("Quốc hội, Luật An ninh mạng số 24/2018/QH14", "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID=129662"),
    ("Thủ tướng Chính phủ, Quyết định 830/QĐ-TTg", "https://vanban.chinhphu.vn/?docid=203332&pageid=27160"),
    ("OWASP Application Security Verification Standard 4.0", "https://owasp.org/www-project-application-security-verification-standard/"),
    ("NIST SP 800-63B Authentication and Authenticator Management", "https://pages.nist.gov/800-63-4/sp800-63b.html"),
    ("FastAPI documentation", "https://fastapi.tiangolo.com/"),
    ("SQLAlchemy ORM documentation", "https://docs.sqlalchemy.org/en/20/orm/"),
    ("Microsoft Windows Service documentation", "https://learn.microsoft.com/windows/win32/services/services"),
]
for i, (title, url) in enumerate(refs, 1):
    add_reference(doc, i, title, url)

add_heading(doc, "PHỤ LỤC A CÁC API CHÍNH", 1)
add_table(doc, ["Nhóm", "Endpoint tiêu biểu", "Chức năng"], [
    ("Auth", "POST /api/auth/login; POST /api/auth/logout", "Phiên phụ huynh"),
    ("Enrollment", "POST /api/enrollment-codes; POST /api/enroll", "Mã ghép và cấp token"),
    ("Agent", "POST /api/heartbeat; GET /api/policy; WS /api/device/ws", "Đồng bộ và lệnh realtime"),
    ("Policy", "GET/PUT /api/children/{id}/policy", "Quota và lịch"),
    ("Filtering", "GET/PUT /api/children/{id}/filtering", "Luật web/app"),
    ("Events", "POST /api/device/activity-events", "Đẩy batch sự kiện"),
    ("Reports", "GET /api/children/{id}/activity-report(.csv/.pdf)", "Báo cáo"),
    ("Deletion", "DELETE /api/children/{id}/activity-data", "Xóa dữ liệu hoạt động"),
    ("Audit", "GET /api/audit; GET /api/device/audit", "Minh bạch hai phía"),
], widths=[2.6, 7.2, 5.8], font_size=9.6)

add_heading(doc, "PHỤ LỤC B TỆP NGUỒN TRỌNG TÂM", 1)
add_table(doc, ["Tệp", "Vai trò"], [
    ("agent/time_control.py", "State machine F1, quota, lịch, idle, grace, monotonic time"),
    ("agent/filtering.py", "So khớp ứng dụng, tên miền, SafeSearch và DNS policy"),
    ("agent/core.py", "Điều phối Service, Tray, report cache và policy"),
    ("agent/remote_sync.py", "Enrollment, token refresh, heartbeat, WebSocket và sync"),
    ("agent/event_queue.py", "SQLite outbox, retention và generation"),
    ("agent/web_protection.py", "Cấu hình/khôi phục DNS và Secure DNS"),
    ("server/app/routes.py", "API danh tính, thiết bị, policy, command và audit"),
    ("server/app/event_routes.py", "Event, report, filtering, deletion và export"),
    ("server/app/security.py", "Argon2id, AES-GCM và token helpers"),
    ("dashboard/static/app.js", "Luồng giao diện phụ huynh và cập nhật realtime"),
    ("docs/PRIVACY.md", "Bảng dữ liệu, mục đích, retention và người xem"),
    ("docs/THREATMODEL.md", "Mô hình STRIDE và giới hạn"),
], widths=[5.2, 10.4])

# Use chapter breaks selectively so the report stays within the required
# 20-30 pages while major transitions still start cleanly.
chapter_breaks = {
    "MỤC LỤC",
    "1 GIỚI THIỆU ĐỀ TÀI",
    "3 KIẾN TRÚC HỆ THỐNG",
    "4 HIỆN THỰC CÁC CHỨC NĂNG",
    "5 PHÂN TÍCH KỸ THUẬT HỆ THỐNG VÀ MẠNG",
    "7 KIỂM THỬ VÀ ĐÁNH GIÁ",
    "9 ĐỐI CHIẾU MỨC ĐỘ HOÀN THÀNH",
    "11 KẾT LUẬN",
}
for paragraph in doc.paragraphs:
    if paragraph.style.name == "Heading 1":
        paragraph.paragraph_format.page_break_before = paragraph.text in chapter_breaks

# Request field refresh in Word.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

doc.core_properties.title = "Báo cáo ứng dụng quản lý máy tính của trẻ em OpenGuard Kids"
doc.core_properties.subject = "Báo cáo đồ án dựa trên yêu cầu và mã nguồn hiện tại"
doc.core_properties.author = "Lê Nhật Thành; Nguyễn Thanh Huyền"
doc.core_properties.keywords = "OpenGuard Kids, Windows agent, FastAPI, parental control, privacy"

doc.save(OUT)
assert sha256(REFERENCE) == EXPECTED_SHA, "Reference template was modified."
print(OUT)
